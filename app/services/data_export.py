"""Выгрузка истории участка в таблицы (pr-087): погода по дням, риск
болезней, спутниковые индексы, обработки и проверки листьев — для отчётов в
Excel, лабораторных работ и анализа в R, Python или QGIS.

XLSX — одна книга, лист на таблицу. CSV — архив, файл на таблицу: в UTF-8 с
меткой порядка байтов (иначе Excel откроет кириллицу как кракозябры) и с
точкой в дробях, чтобы файл без настройки читали pandas и R."""

import csv
import io
import uuid
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from enum import StrEnum
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.diagnosis import DiseaseResult, LeafImage
from app.models.satellite import IndexType
from app.models.treatment import Treatment
from app.models.weather import DiseaseRisk, WeatherData
from app.services.agromodels import LOCAL_TZ, HourlyWeather, daily_weather
from app.services.vegetation_history import vegetation_history

# Не больше трёх лет за раз: почасовая погода участка за такой срок — около
# 26 тысяч строк, больше в одном запросе собирать незачем
MAX_DAYS = 3 * 366


class DataFormat(StrEnum):
    """Формат выгрузки истории участка: книга XLSX или архив CSV."""

    xlsx = "xlsx"
    csv = "csv"


Cell = str | int | float | date | datetime | None


@dataclass(frozen=True)
class Table:
    name: str
    filename: str
    header: list[str]
    rows: list[list[Cell]]


def _utc(day: date) -> datetime:
    return datetime.combine(day, time(), LOCAL_TZ).astimezone(UTC)


async def collect_tables(
    session: AsyncSession, vineyard_id: uuid.UUID, start: date, end: date
) -> list[Table]:
    """Таблицы участка за местные сутки с start по end включительно."""
    since, until = _utc(start), _utc(end + timedelta(days=1))

    hours = await session.execute(
        select(
            WeatherData.observed_at,
            WeatherData.temperature_c,
            WeatherData.relative_humidity,
            WeatherData.precipitation_mm,
        )
        .where(
            WeatherData.vineyard_id == vineyard_id,
            WeatherData.observed_at >= since,
            WeatherData.observed_at < until,
            # Только факт: прогноз на будущие часы — не история
            WeatherData.is_forecast.is_(False),
        )
        .order_by(WeatherData.observed_at)
    )
    weather = Table(
        "Погода по дням",
        "weather_daily.csv",
        [
            "date",
            "mean_temp_c",
            "min_temp_c",
            "max_temp_c",
            "precipitation_mm",
            "mean_humidity",
            "mild_hours",
        ],
        [
            [
                day.day,
                day.mean_temp,
                day.min_temp,
                day.max_temp,
                day.precipitation_mm,
                day.mean_humidity,
                day.longest_mild_run,
            ]
            for day in daily_weather(HourlyWeather(*row) for row in hours)
        ],
    )

    risks = await session.scalars(
        select(DiseaseRisk)
        .where(
            DiseaseRisk.vineyard_id == vineyard_id,
            DiseaseRisk.target_date >= start,
            DiseaseRisk.target_date <= end,
        )
        .order_by(DiseaseRisk.target_date, DiseaseRisk.disease, DiseaseRisk.method)
    )
    risk = Table(
        "Риск",
        "disease_risk.csv",
        ["date", "disease", "method", "score", "level"],
        [
            [r.target_date, r.disease.value, r.method, r.score, r.level.value]
            for r in risks
        ],
    )

    index_rows: list[list[Cell]] = []
    for scan in await vegetation_history(session, vineyard_id):
        if not since <= scan.acquired_at < until:
            continue
        values = {index.index_type: index.mean_value for index in scan.indices}
        index_rows.append(
            [
                scan.acquired_at,
                scan.platform,
                values.get(IndexType.NDVI),
                values.get(IndexType.NDRE),
                scan.parcel_cloud_cover,
            ]
        )
    indices = Table(
        "Индексы",
        "vegetation_indices.csv",
        ["acquired_at", "platform", "ndvi_mean", "ndre_mean", "parcel_cloud_cover"],
        index_rows,
    )

    treatments = await session.scalars(
        select(Treatment)
        .where(
            Treatment.vineyard_id == vineyard_id,
            Treatment.applied_on >= start,
            Treatment.applied_on <= end,
        )
        .order_by(Treatment.applied_on)
    )
    treatment = Table(
        "Обработки",
        "treatments.csv",
        ["date", "fungicide", "product"],
        [[t.applied_on, t.fungicide.value, t.product] for t in treatments],
    )

    checks = await session.execute(
        select(LeafImage.created_at, DiseaseResult)
        .join(DiseaseResult, DiseaseResult.image_id == LeafImage.id)
        .where(
            LeafImage.vineyard_id == vineyard_id,
            LeafImage.created_at >= since,
            LeafImage.created_at < until,
        )
        .order_by(LeafImage.created_at)
    )
    diagnosis = Table(
        "Проверки листьев",
        "leaf_checks.csv",
        ["created_at", "predicted", "confidence", "severity", "expert"],
        [
            [
                created,
                result.predicted_class,
                result.confidence,
                result.severity,
                result.expert_class,
            ]
            for created, result in checks
        ],
    )
    return [weather, risk, indices, treatment, diagnosis]


def _csv_value(value: Cell) -> Any:
    if isinstance(value, datetime):
        return value.astimezone(UTC).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else value


def to_csv_zip(tables: Sequence[Table]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for table in tables:
            text = io.StringIO()
            writer = csv.writer(text, lineterminator="\n")
            writer.writerow(table.header)
            writer.writerows([_csv_value(cell) for cell in row] for row in table.rows)
            archive.writestr(table.filename, "﻿" + text.getvalue())
    return buffer.getvalue()


def _xlsx_value(value: Cell) -> Any:
    # Excel не хранит часовой пояс: время пишется в UTC, о чём сказано в заголовке
    if isinstance(value, datetime):
        return value.astimezone(UTC).replace(tzinfo=None)
    return value


def to_xlsx(tables: Sequence[Table]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for table in tables:
        sheet = workbook.create_sheet(table.name)
        header = [
            f"{name} (UTC)" if name.endswith("_at") else name for name in table.header
        ]
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in table.rows:
            sheet.append([_xlsx_value(cell) for cell in row])
        # Шапка на месте при прокрутке, столбцы — по ширине заголовка
        sheet.freeze_panes = "A2"
        for column, name in enumerate(header, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = max(
                12, len(name) + 2
            )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
