import csv
import io
import zipfile
from collections.abc import Callable
from datetime import UTC, date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.data_export import Table, to_csv_zip, to_xlsx

RunSql = Callable[[str], list[tuple[object, ...]]]

TABLES = [
    Table(
        "Погода по дням",
        "weather_daily.csv",
        ["date", "mean_temp_c", "note"],
        [[date(2026, 6, 1), 18.5, "Криково"], [date(2026, 6, 2), 20.25, None]],
    ),
    Table(
        "Индексы",
        "vegetation_indices.csv",
        ["acquired_at", "ndvi_mean"],
        [[datetime(2026, 6, 1, 12, 17, tzinfo=UTC), 0.71]],
    ),
]


def test_csv_archive_opens_in_excel_and_pandas() -> None:
    archive = zipfile.ZipFile(io.BytesIO(to_csv_zip(TABLES)))

    assert archive.namelist() == ["weather_daily.csv", "vegetation_indices.csv"]
    raw = archive.read("weather_daily.csv").decode("utf-8")
    # Метка порядка байтов: Excel узнаёт UTF-8 и не ломает кириллицу
    assert raw.startswith("﻿")
    rows = list(csv.reader(io.StringIO(raw.removeprefix("﻿"))))
    assert rows == [
        ["date", "mean_temp_c", "note"],
        ["2026-06-01", "18.5", "Криково"],
        ["2026-06-02", "20.25", ""],
    ]
    indices = archive.read("vegetation_indices.csv").decode("utf-8-sig")
    assert "2026-06-01T12:17:00+00:00,0.71" in indices


def test_workbook_has_sheet_per_table() -> None:
    workbook = load_workbook(io.BytesIO(to_xlsx(TABLES)))

    assert workbook.sheetnames == ["Погода по дням", "Индексы"]
    weather = workbook["Погода по дням"]
    assert [cell.value for cell in weather[1]] == ["date", "mean_temp_c", "note"]
    assert weather.freeze_panes == "A2"
    assert weather["B3"].value == 20.25
    indices = workbook["Индексы"]
    # В Excel нет часовых поясов: время в UTC, и это сказано в заголовке
    assert indices["A1"].value == "acquired_at (UTC)"
    assert indices["A2"].value == datetime(2026, 6, 1, 12, 17)


def login(client: TestClient, email: str = "grower@example.md") -> dict[str, str]:
    credentials = {"email": email, "password": "secret123"}
    client.post("/api/v1/auth/register", json={**credentials, "full_name": "Ион"})
    token = client.post("/api/v1/auth/login", json=credentials).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.integration
def test_vineyard_history_export(api_client: TestClient, run_sql: RunSql) -> None:
    headers = login(api_client)
    ring = [[28.86, 47.14], [28.862, 47.14], [28.862, 47.141], [28.86, 47.14]]
    vineyard = api_client.post(
        "/api/v1/vineyards",
        headers=headers,
        json={
            "name": "Криково",
            "geometry": {"type": "Polygon", "coordinates": [ring]},
        },
    ).json()["id"]
    owner = api_client.get("/api/v1/users/me", headers=headers).json()["id"]
    # Сутки 1 июня по Кишинёву (UTC+3): с 21:00 31 мая по UTC, по 0,5 мм в час
    run_sql(
        "INSERT INTO weather_data (vineyard_id, observed_at, temperature_c, "
        "relative_humidity, precipitation_mm, is_forecast) SELECT "
        f"'{vineyard}', timestamp with time zone '2026-05-31 21:00+00' + "
        "make_interval(hours => h), 20, 70, 0.5, false FROM generate_series(0, 23) h"
    )
    # Прогноз в выгрузку истории не попадает
    run_sql(
        "INSERT INTO weather_data (vineyard_id, observed_at, temperature_c, "
        "relative_humidity, precipitation_mm, is_forecast) VALUES "
        f"('{vineyard}', '2026-06-05 12:00+00', 30, 50, 0, true)"
    )
    run_sql(
        "INSERT INTO disease_risks (vineyard_id, disease, target_date, method, score, "
        f"level) VALUES ('{vineyard}', 'downy_mildew', '2026-06-01', 'three_tens', "
        "1, 'high'), "
        f"('{vineyard}', 'downy_mildew', '2025-06-01', 'three_tens', 1, 'high')"
    )
    run_sql(
        "WITH scan AS (INSERT INTO satellite_scans (vineyard_id, product_id, "
        "acquired_at, platform, cloud_cover, parcel_cloud_cover, status) VALUES "
        f"('{vineyard}', 'S2_export', '2026-06-01T09:17:00+00', 'sentinel-2a', 0, 3, "
        f"'completed'), ('{vineyard}', 'S2_may', '2026-05-01T09:17:00+00', "
        "'sentinel-2a', 0, 0, 'completed') RETURNING id) INSERT INTO "
        "vegetation_indices (scan_id, index_type, mean_value, min_value, max_value, "
        "std_value) SELECT id, 'ndvi', 0.7, 0.6, 0.8, 0.05 FROM scan"
    )
    run_sql(
        "INSERT INTO treatments (vineyard_id, applied_on, fungicide, product) VALUES "
        f"('{vineyard}', '2026-06-02', 'contact', 'Купроксат')"
    )
    run_sql(
        "WITH image AS (INSERT INTO leaf_images (owner_id, vineyard_id, storage_key, "
        f"content_type, size_bytes, width, height, created_at) VALUES ('{owner}', "
        f"'{vineyard}', 'leaf/export.jpg', 'image/jpeg', 10, 224, 224, "
        "'2026-06-03T08:00:00+00') RETURNING id) "
        "INSERT INTO disease_results (image_id, model_name, model_version, status, "
        "predicted_class, confidence) SELECT id, 'efficientnet_b0', 'v1', "
        "'completed', 'downy_mildew', 0.91 FROM image"
    )
    url = f"/api/v1/vineyards/{vineyard}/data"
    period = {"start": "2026-06-01", "end": "2026-06-30"}

    response = api_client.get(url, headers=headers, params=period)

    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        'attachment; filename="vitiguard-data-2026-06-01-2026-06-30.xlsx"'
    )
    workbook = load_workbook(io.BytesIO(response.content))
    weather = list(workbook["Погода по дням"].values)
    assert weather[1][:5] == (datetime(2026, 6, 1), 20, 20, 20, 12)
    assert len(weather) == 2
    assert [row[:3] for row in list(workbook["Риск"].values)[1:]] == [
        (datetime(2026, 6, 1), "downy_mildew", "three_tens")
    ]
    # Майский снимок — вне периода
    indices = list(workbook["Индексы"].values)
    assert len(indices) == 2
    assert indices[1][1:3] == ("sentinel-2a", 0.7)
    assert list(workbook["Обработки"].values)[1][1:] == ("contact", "Купроксат")
    assert list(workbook["Проверки листьев"].values)[1][1:3] == ("downy_mildew", 0.91)

    archive = api_client.get(url, headers=headers, params={**period, "format": "csv"})
    assert archive.headers["content-type"] == "application/zip"
    assert len(zipfile.ZipFile(io.BytesIO(archive.content)).namelist()) == 5

    reversed_period = {"start": "2026-06-30", "end": "2026-06-01"}
    assert (
        api_client.get(url, headers=headers, params=reversed_period).status_code == 400
    )
    too_long = {"start": "2020-01-01", "end": "2026-06-01"}
    assert api_client.get(url, headers=headers, params=too_long).status_code == 400
    stranger = login(api_client, "stranger@example.md")
    assert api_client.get(url, headers=stranger).status_code == 404
    assert api_client.get(url).status_code == 401


@pytest.mark.integration
def test_default_period_is_current_season(
    api_client: TestClient, run_sql: RunSql
) -> None:
    from app.services.risk import local_today, season_year

    headers = login(api_client)
    ring = [[28.86, 47.14], [28.862, 47.14], [28.862, 47.141], [28.86, 47.14]]
    vineyard = api_client.post(
        "/api/v1/vineyards",
        headers=headers,
        json={
            "name": "Криково",
            "geometry": {"type": "Polygon", "coordinates": [ring]},
        },
    ).json()["id"]
    today = local_today()
    old = date(season_year(today), 3, 1) - timedelta(days=1)
    run_sql(
        "INSERT INTO treatments (vineyard_id, applied_on, fungicide) VALUES "
        f"('{vineyard}', '{old}', 'sulfur'), ('{vineyard}', '{today}', 'contact')"
    )

    response = api_client.get(f"/api/v1/vineyards/{vineyard}/data", headers=headers)

    rows = list(load_workbook(io.BytesIO(response.content))["Обработки"].values)[1:]
    assert [row[1] for row in rows] == ["contact"]
